"""Construit le corpus de Nexalis Conseil dans data/raw/ à partir des sources Markdown.

Chaque document est rédigé une fois en Markdown dans data/sources/ (avec un en-tête YAML de
métadonnées), puis rendu dans son format cible pour tester le parsing : DOCX, PDF texte,
PDF scanné (image seule, pour l'OCR) ou Markdown. La convention Syntec est récupérée depuis
le jeu de test public lyon-nlp/mteb-fr-retrieval-syntec-s2p.

Usage : python -m scripts.build_corpus
"""

import io
import json
import random
import re
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[1]
SOURCES_DIR = ROOT / "data" / "sources"
RAW_DIR = ROOT / "data" / "raw"
EVAL_DIR = ROOT / "eval" / "syntec"
SYNTEC_DATASET = "lyon-nlp/mteb-fr-retrieval-syntec-s2p"
SYNTEC_FILE = "convention_collective_syntec.md"

FOOTER_PREFIX = "Document interne fictif (projet pédagogique Ynov). Sources : "


# --------------------------------------------------------------------------- lecture Markdown


def read_source(path: Path) -> tuple[dict, str]:
    """Sépare l'en-tête YAML et le corps Markdown."""
    text = path.read_text(encoding="utf-8")
    _, front, body = text.split("---", 2)
    return yaml.safe_load(front), body.strip()


def parse_blocks(body: str) -> list[tuple[str, object]]:
    """Découpe un Markdown simple en blocs : titre, paragraphe, liste, liste numérotée, tableau."""
    blocks: list[tuple[str, object]] = []
    lines = body.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if m := re.match(r"^(#{1,6})\s+(.*)", line):
            blocks.append(("heading", (len(m.group(1)), m.group(2))))
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
        elif re.match(r"^\s*[-*]\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("bullets", items))
        elif re.match(r"^\s*\d+\.\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+\.\s+", "", lines[i]).strip())
                i += 1
            blocks.append(("numbers", items))
        else:
            para = [line.strip()]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(
                r"^(#|\||\s*[-*]\s|\s*\d+\.\s)", lines[i]
            ):
                para.append(lines[i].strip())
                i += 1
            blocks.append(("paragraph", " ".join(para)))
    return blocks


def inline_runs(text: str) -> list[tuple[str, bool]]:
    """Découpe un texte en segments (texte, gras ?) selon la syntaxe **gras**. Retire les `code`."""
    text = text.replace("`", "")
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") else (p, False) for p in parts if p]


# --------------------------------------------------------------------------- rendu DOCX


def render_docx(meta: dict, blocks: list, out: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.core_properties.title = meta["titre"]
    doc.core_properties.author = "Nexalis Conseil"
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def add_runs(paragraph, text):
        for chunk, bold in inline_runs(text):
            paragraph.add_run(chunk).bold = bold

    for kind, value in blocks:
        if kind == "heading":
            level, text = value
            doc.add_heading(text.replace("**", ""), level=min(level - 1, 3) if level > 1 else 0)
        elif kind == "paragraph":
            add_runs(doc.add_paragraph(), value)
        elif kind in ("bullets", "numbers"):
            style_name = "List Bullet" if kind == "bullets" else "List Number"
            for item in value:
                add_runs(doc.add_paragraph(style=style_name), item)
        elif kind == "table":
            table = doc.add_table(rows=len(value), cols=len(value[0]))
            table.style = "Table Grid"
            for r, row in enumerate(value):
                for c, cell in enumerate(row):
                    paragraph = table.cell(r, c).paragraphs[0]
                    for chunk, bold in inline_runs(cell):
                        paragraph.add_run(chunk).bold = bold or r == 0

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = FOOTER_PREFIX + " ".join(meta["sources"].split())
    footer.runs[0].font.size = Pt(8)
    doc.save(out)


# --------------------------------------------------------------------------- rendu PDF


def _register_font() -> str:
    """Police TrueType embarquée (accents, tirets) : DejaVu Sans fournie avec matplotlib."""
    from matplotlib import get_data_path
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_dir = Path(get_data_path()) / "fonts" / "ttf"
    pdfmetrics.registerFont(TTFont("DejaVu", str(font_dir / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(font_dir / "DejaVuSans-Bold.ttf")))
    addMapping("DejaVu", 0, 0, "DejaVu")
    addMapping("DejaVu", 1, 0, "DejaVu-Bold")
    addMapping("DejaVu", 0, 1, "DejaVu")
    addMapping("DejaVu", 1, 1, "DejaVu-Bold")
    return "DejaVu"


def _xml(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return "".join(f"<b>{t}</b>" if bold else t for t, bold in inline_runs(text))


def render_pdf(meta: dict, blocks: list, out: Path | io.BytesIO, font_size: float = 10) -> None:
    from reportlab import rl_config
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    rl_config.invariant = 1  # PDF reproductible (pas de date de création) : empreinte stable
    font = _register_font()
    body = ParagraphStyle("body", fontName=font, fontSize=font_size, leading=font_size * 1.4,
                          spaceAfter=6)
    cell = ParagraphStyle("cell", parent=body, fontSize=font_size - 1, leading=font_size * 1.2,
                          spaceAfter=0)
    headings = {
        1: ParagraphStyle("h1", parent=body, fontName=f"{font}-Bold", fontSize=18, leading=22,
                          spaceAfter=12),
        2: ParagraphStyle("h2", parent=body, fontName=f"{font}-Bold", fontSize=13.5, leading=17,
                          spaceBefore=10, spaceAfter=6),
        3: ParagraphStyle("h3", parent=body, fontName=f"{font}-Bold", fontSize=11.5, leading=15,
                          spaceBefore=8, spaceAfter=4),
    }
    footer_text = FOOTER_PREFIX + " ".join(meta["sources"].split())

    def on_page(canvas, doc_template):
        canvas.saveState()
        footer_style = ParagraphStyle("footer", fontName=font, fontSize=6.5, leading=8,
                                      textColor=colors.grey)
        footer = Paragraph(_xml(footer_text), footer_style)
        width = A4[0] - 4 * cm
        _, height = footer.wrap(width, 3 * cm)
        footer.drawOn(canvas, 2 * cm, 1.0 * cm)
        canvas.setFont(font, 8)
        canvas.drawRightString(A4[0] - 2 * cm, 1.0 * cm + height + 4,
                               f"{meta['titre']} — page {doc_template.page}")
        canvas.restoreState()

    story = []
    for kind, value in blocks:
        if kind == "heading":
            level, text = value
            story.append(Paragraph(_xml(text), headings[min(level, 3)]))
        elif kind == "paragraph":
            story.append(Paragraph(_xml(value), body))
        elif kind in ("bullets", "numbers"):
            items = [ListItem(Paragraph(_xml(item), body), leftIndent=12) for item in value]
            story.append(ListFlowable(items, bulletType="bullet" if kind == "bullets" else "1",
                                      start="•" if kind == "bullets" else 1,
                                      bulletFontName=font, bulletFontSize=font_size))
        elif kind == "table":
            data = [[Paragraph(_xml(c), cell) for c in row] for row in value]
            table = Table(data, repeatRows=1, hAlign="LEFT",
                          colWidths=[(A4[0] - 4 * cm) / len(value[0])] * len(value[0]))
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dfe7f2")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
            story.append(Spacer(1, 8))

    template = SimpleDocTemplate(str(out) if isinstance(out, Path) else out, pagesize=A4,
                                 leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                                 bottomMargin=2.8 * cm, title=meta["titre"],
                                 author="Nexalis Conseil")
    template.build(story, onFirstPage=on_page, onLaterPages=on_page)


def render_scanned_pdf(meta: dict, blocks: list, out: Path, dpi: int = 200) -> None:
    """PDF texte rasterisé en images légèrement bruitées et inclinées : un vrai « scan »."""
    import pypdfium2 as pdfium
    from PIL import Image, ImageFilter

    buffer = io.BytesIO()
    render_pdf(meta, blocks, buffer, font_size=11)
    pdf = pdfium.PdfDocument(buffer.getvalue())
    rng = random.Random(42)
    pages = []
    for page in pdf:
        image = page.render(scale=dpi / 72).to_pil().convert("L")
        image = image.rotate(rng.uniform(-0.8, 0.8), fillcolor=255, expand=False)
        noise = Image.effect_noise(image.size, 18).filter(ImageFilter.GaussianBlur(0.6))
        image = Image.blend(image, noise, 0.06).filter(ImageFilter.GaussianBlur(0.4))
        pages.append(image.convert("RGB"))
    pages[0].save(out, save_all=True, append_images=pages[1:], resolution=dpi)


def render_markdown(meta: dict, body: str, out: Path) -> None:
    footer = FOOTER_PREFIX + " ".join(meta["sources"].split())
    out.write_text(f"{body}\n\n---\n\n_{footer}_\n", encoding="utf-8")


# --------------------------------------------------------------------------- convention Syntec


def build_syntec() -> dict:
    """Écrit la convention Syntec (90 articles) en Markdown et les 100 questions annotées."""
    from datasets import load_dataset

    docs = load_dataset(SYNTEC_DATASET, "documents", split="test")
    queries = load_dataset(SYNTEC_DATASET, "queries", split="test")

    lines = [
        "# Convention collective nationale Syntec (IDCC 1486)",
        "",
        "Convention collective nationale des bureaux d'études techniques, des cabinets "
        "d'ingénieurs-conseils et des sociétés de conseils, applicable à Nexalis Conseil.",
        "",
    ]
    current_section = None
    articles = []
    # Le jeu est trié par identifiant (article-1, article-10, article-11...) : on rétablit l'ordre
    # du texte, préambules d'abord puis articles par numéro.
    def order(row):
        num = row["id"].removeprefix("article-")
        return (1, int(num)) if num.isdigit() else (0, 0)

    for row in sorted(docs, key=order):
        title = " ".join(row["title"].split())
        if row["section"] != current_section:
            current_section = row["section"]
            lines += [f"## {current_section}", ""]
        lines += [f"### {title}", "", row["content"].strip(), ""]
        articles.append({"id": row["id"], "title": title, "section": row["section"],
                         "url": row["url"]})

    meta = {
        "titre": "Convention collective nationale Syntec",
        "fichier": SYNTEC_FILE,
        "format": "md",
        "categorie": "RH",
        "version": "texte du jeu MTEB-French (2024)",
        "date_effet": None,
        "statut": "en_vigueur",
        "sources": "Texte de la convention Syntec (IDCC 1486, public sur Légifrance) tel que "
        f"publié dans le jeu de test {SYNTEC_DATASET} "
        "(Wikit AI, MTEB-French, Ciancone et al. 2024).",
    }
    (RAW_DIR / SYNTEC_FILE).write_text("\n".join(lines).strip() + "\n", encoding="utf-8")

    EVAL_DIR.mkdir(parents=True, exist_ok=True)
    with (EVAL_DIR / "queries.jsonl").open("w", encoding="utf-8") as f:
        for row in queries:
            f.write(json.dumps({"question": row["Question"], "article": row["Article"]},
                               ensure_ascii=False) + "\n")
    (EVAL_DIR / "articles.json").write_text(
        json.dumps(articles, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    return meta


def write_sources(metadata: dict) -> None:
    """data/SOURCES.md : la source réelle de chaque document (hors data/raw, non indexé)."""
    lines = [
        "# Sources du corpus Nexalis Conseil",
        "",
        "Nexalis Conseil est une entreprise fictive. Chaque document interne s'appuie sur une "
        "source publique réelle ; les règles propres à Nexalis (montants, plafonds, délais) sont "
        "inventées. Fichier généré par `python -m scripts.build_corpus`.",
        "",
        "| Document | Format | Version | Statut | Sources |",
        "| --- | --- | --- | --- | --- |",
    ]
    for name, meta in sorted(metadata.items()):
        sources = " ".join(str(meta["sources"]).split())
        lines.append(f"| {meta['titre']} (`{name}`) | {meta['format']} | {meta['version']} | "
                     f"{meta['statut']} | {sources} |")
    (RAW_DIR.parent / "SOURCES.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


# --------------------------------------------------------------------------- main


def render_source(path: Path) -> tuple[str, dict]:
    """Rend un document source dans son format cible ; renvoie (fichier, métadonnées)."""
    meta, body = read_source(path)
    blocks = parse_blocks(body)
    out = RAW_DIR / meta["fichier"]
    if meta["format"] == "docx":
        render_docx(meta, blocks, out)
    elif meta["format"] == "pdf":
        render_pdf(meta, blocks, out)
    elif meta["format"] == "pdf_scan":
        render_scanned_pdf(meta, blocks, out)
    else:
        render_markdown(meta, body, out)
    print(f"  {meta['format']:<8} -> {out.relative_to(ROOT)}")
    return meta["fichier"], {k: v for k, v in meta.items() if k != "fichier"}


def build_only(names: list[str]) -> None:
    """Rend seulement les sources indiquées et complète metadata.json : les autres documents
    restent identiques octet pour octet, donc l'ingestion incrémentale ne les retraite pas."""
    metadata_path = RAW_DIR / "metadata.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    for name in names:
        path = SOURCES_DIR / (name if name.endswith(".md") else f"{name}.md")
        fichier, meta = render_source(path)
        metadata[fichier] = meta
    metadata_path.write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    write_sources(metadata)
    print(f"{len(names)} document(s) rendu(s), {len(metadata)} documents dans metadata.json")


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Construit le corpus dans data/raw.")
    parser.add_argument("--only", nargs="+", metavar="SOURCE",
                        help="ne rendre que ces sources de data/sources (sans extension)")
    args = parser.parse_args()
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    if args.only:
        build_only(args.only)
        return

    metadata = {}
    for path in sorted(SOURCES_DIR.glob("*.md")):
        meta, body = read_source(path)
        blocks = parse_blocks(body)
        out = RAW_DIR / meta["fichier"]
        if meta["format"] == "docx":
            render_docx(meta, blocks, out)
        elif meta["format"] == "pdf":
            render_pdf(meta, blocks, out)
        elif meta["format"] == "pdf_scan":
            render_scanned_pdf(meta, blocks, out)
        else:
            render_markdown(meta, body, out)
        metadata[meta["fichier"]] = {k: v for k, v in meta.items() if k != "fichier"}
        print(f"  {meta['format']:<8} -> {out.relative_to(ROOT)}")

    syntec = build_syntec()
    metadata[SYNTEC_FILE] = {k: v for k, v in syntec.items() if k != "fichier"}
    print(f"  {'md':<8} -> data/raw/{SYNTEC_FILE} (+ eval/syntec/)")

    (RAW_DIR / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    write_sources(metadata)
    print(f"{len(metadata)} documents, métadonnées dans data/raw/metadata.json")


if __name__ == "__main__":
    main()
